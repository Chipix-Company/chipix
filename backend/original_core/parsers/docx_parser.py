"""
DOCX specification parser.

Extracts text and tables from Word documents.
"""

from __future__ import annotations

from pathlib import Path

from core.logger import get_logger

logger = get_logger("Parser.DOCX")


def parse_docx(file_path: str | Path) -> str:
    """Extract all text and tables from a DOCX file."""
    file_path = Path(file_path)
    if not file_path.exists():
        raise FileNotFoundError(f"DOCX file not found: {file_path}")

    try:
        from docx import Document
    except ImportError:
        raise ImportError("python-docx is not installed. Install: pip install python-docx")

    doc = Document(str(file_path))
    text_parts = []

    # Extract paragraphs
    for para in doc.paragraphs:
        if para.text.strip():
            # Preserve heading hierarchy
            if para.style.name.startswith("Heading"):
                level = para.style.name.replace("Heading ", "").strip()
                prefix = "#" * int(level) if level.isdigit() else "#"
                text_parts.append(f"{prefix} {para.text.strip()}")
            else:
                text_parts.append(para.text.strip())

    # Extract tables
    for table in doc.tables:
        table_rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            table_rows.append(" | ".join(cells))
        if table_rows:
            text_parts.append("\n".join(table_rows))

    text = "\n\n".join(text_parts)
    logger.info("Parsed DOCX: %d chars, %d paragraphs, %d tables",
                 len(text), len(doc.paragraphs), len(doc.tables))
    return text
